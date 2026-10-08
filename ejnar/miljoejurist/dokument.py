"""Læs et uploadet dokument (PDF, Word, HTML eller tekst) og del det i sætninger.

Sætningerne beholder deres placering, så rapportens citater kan vises og
kontrolleres mod dokumentet.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass


@dataclass
class Sætning:
    nr: int
    tekst: str
    start: int


@dataclass
class Dokument:
    navn: str
    tekst: str
    sætninger: list[Sætning]

    @property
    def ord(self) -> int:
        return len(self.tekst.split())


def _pdf(data: bytes) -> str:
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            return "\n\n".join((p.extract_text() or "") for p in pdf.pages)
    except ImportError:
        from pypdf import PdfReader
        return "\n\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)


def _docx(data: bytes) -> str:
    import docx  # python-docx
    d = docx.Document(io.BytesIO(data))
    dele = [p.text for p in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            dele.append(" | ".join(c.text.strip() for c in r.cells))
    return "\n".join(dele)


def _html(data: bytes) -> str:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(data, "html.parser")
    for t in soup(["script", "style"]):
        t.decompose()
    return soup.get_text("\n")


def tekst_fra_fil(navn: str, data: bytes) -> str:
    n = (navn or "").lower()
    if n.endswith(".pdf") or data[:4] == b"%PDF":
        t = _pdf(data)
    elif n.endswith(".docx"):
        t = _docx(data)
    elif n.endswith((".html", ".htm")):
        t = _html(data)
    else:
        for enc in ("utf-8", "cp1252", "latin-1"):
            try:
                t = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
    t = t.replace("\xa0", " ").replace("\r", "")
    t = re.sub(r"(\w)-\n(\w)", r"\1\2", t)          # orddeling ved linjeskift
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


_SÆT = re.compile(r"(?<=[.!?:;])\s+(?=[A-ZÆØÅ0-9•\-–(\"»“])|\n\s*\n|\n(?=\s*[•\-–]\s)")
_FORK = re.compile(r"(?:\b(?:jf|nr|stk|bl\.a|f\.eks|ca|mv|pkt|m\.fl|evt|inkl|ift|vedr|kap|afs)\.|\b[a-z]\.)$", re.I)


def sætninger(tekst: str) -> list[Sætning]:
    ud, pos, buf, buf_start = [], 0, "", 0
    for m in _SÆT.finditer(tekst):
        stykke = tekst[pos:m.start()]
        if not buf:
            buf_start = pos
        buf += stykke
        pos = m.end()
        if _FORK.search(buf.strip()):   # "jf." o.l. er ikke sætningsslut
            buf += " "
            continue
        if buf.strip():
            ud.append(Sætning(len(ud), re.sub(r"\s+", " ", buf).strip(), buf_start))
        buf = ""
    rest = buf + tekst[pos:]
    if rest.strip():
        ud.append(Sætning(len(ud), re.sub(r"\s+", " ", rest).strip(), buf_start if buf else pos))
    return [s for s in ud if len(s.tekst) > 2]


def læs(navn: str, data: bytes) -> Dokument:
    t = tekst_fra_fil(navn, data)
    return Dokument(navn, t, sætninger(t))


def fra_tekst(tekst: str, navn: str = "indsat tekst") -> Dokument:
    return læs(navn, tekst.encode("utf-8"))
